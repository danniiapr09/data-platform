import io
import os

os.environ["DATABASE_URL"] = "sqlite:///./test.db"

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

ROWS = "\n".join(f"{2018 + i},{1000 + 40 * i},{'=HYPERLINK(1)' if i == 0 else 'ok'}" for i in range(6))
CSV = "tahun,siswa,catatan\n" + ROWS + "\n"


def setup(c, forecast=True):
    did = c.post("/datasets/upload", files={"file": ("d.csv", CSV, "text/csv")}).json()["dataset_id"]
    c.post(f"/datasets/{did}/clean")
    if forecast:
        assert c.post(f"/datasets/{did}/forecast", json={"nilai": "siswa"}).status_code == 200
    return did


def test_detail_report_history():
    with TestClient(app) as c:
        did = setup(c)
        d = c.get(f"/datasets/{did}").json()
        assert d["baris_mentah"] == 6 and d["versi_bersih"] == 1 and d["ada_forecast"] is True
        r = c.get(f"/datasets/{did}/clean-report").json()
        assert r["baris_awal"] == 6 and "kelengkapan_persen" in r
        h = c.get(f"/datasets/{did}/history").json()
        assert [x["langkah"] for x in h] == ["upload", "cleaning", "forecasting"]
        assert c.get("/datasets/9999").status_code == 404
        assert c.get("/datasets/9999/clean-report").status_code == 404


def test_report_404_before_clean():
    with TestClient(app) as c:
        did = c.post("/datasets/upload", files={"file": ("d.csv", CSV, "text/csv")}).json()["dataset_id"]
        assert c.get(f"/datasets/{did}/clean-report").status_code == 404
        assert c.get(f"/datasets/{did}/export/xlsx").status_code == 400


def test_export_xlsx_sheets_and_formula_guard():
    with TestClient(app) as c:
        did = setup(c)
        r = c.get(f"/datasets/{did}/export/xlsx")
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Ringkasan", "Data Bersih", "Anomali", "Forecast"]
        cells = [row[2].value for row in wb["Data Bersih"].iter_rows(min_row=2)]
        assert cells[0] == "'=HYPERLINK(1)"           # rumus dinetralkan
        ring = {row[0].value: row[1].value for row in wb["Ringkasan"].iter_rows(min_row=2)}
        assert ring["Baris awal"] == 6 and "Catatan prediksi" in ring
        assert wb["Forecast"].max_row == 4              # header + 3 prediksi


def test_export_without_forecast_and_csv():
    with TestClient(app) as c:
        did = setup(c, forecast=False)
        wb = load_workbook(io.BytesIO(c.get(f"/datasets/{did}/export/xlsx").content))
        assert "Forecast" not in wb.sheetnames
        assert c.get(f"/datasets/{did}/export/forecast.csv").status_code == 404
        c.post(f"/datasets/{did}/forecast", json={"nilai": "siswa", "horizon": 2})
        r = c.get(f"/datasets/{did}/export/forecast.csv")
        assert r.status_code == 200 and r.content.decode("utf-8-sig").splitlines()[0].startswith("indikator")
        assert len(r.content.decode("utf-8-sig").strip().splitlines()) == 3